import csv
import io
import json
import os
import re
from datetime import datetime, timedelta, timezone
from functools import wraps
from uuid import uuid4

from flask import Flask, jsonify, request, send_file
from flask_bcrypt import Bcrypt
from flask_cors import CORS
from flask_jwt_extended import (
    JWTManager,
    create_access_token,
    get_jwt,
    get_jwt_identity,
    jwt_required,
)

from models import get_db


UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")
ALLOWED_UPLOAD_EXTENSIONS = {".json", ".txt", ".pdf", ".docx", ".doc"}
REQUIRED_COLUMNS = {"StudentID", "StudentName"}
GRADE_RULES = (("A", 90), ("B", 75), ("C", 50))


app = Flask(__name__)
CORS(app)
bcrypt = Bcrypt(app)
jwt = JWTManager(app)

app.config["JWT_SECRET_KEY"] = os.getenv("JWT_SECRET_KEY", "change-this-secret-in-production")
app.config["JWT_ACCESS_TOKEN_EXPIRES"] = timedelta(hours=4)

db = get_db()
users_collection = db["users"]
students_collection = db["students"]
classes_collection = db["classes"]

os.makedirs(UPLOAD_DIR, exist_ok=True)


def utc_now():
    return datetime.now(timezone.utc)


def to_iso(dt):
    if isinstance(dt, datetime):
        return dt.astimezone(timezone.utc).isoformat()
    return dt


def is_valid_email(email: str) -> bool:
    return isinstance(email, str) and "@" in email and "." in email


def role_required(required_role: str):
    def decorator(fn):
        @wraps(fn)
        @jwt_required()
        def wrapper(*args, **kwargs):
            claims = get_jwt()
            if claims.get("role") != required_role:
                return jsonify({"error": "Forbidden: insufficient role permissions"}), 403
            return fn(*args, **kwargs)

        return wrapper

    return decorator


def build_class_id():
    return f"class_{utc_now().strftime('%Y%m%d%H%M%S')}_{uuid4().hex[:8]}"


def normalize_class_name(class_name: str) -> str:
    cleaned = str(class_name or "").strip()
    if cleaned:
        return cleaned
    return f"Class Upload {utc_now().strftime('%d %b %Y %H:%M')}"


def normalize_row_keys(row):
    normalized = {}
    for key, value in (row or {}).items():
        if key is None:
            continue
        normalized[str(key).strip()] = value
    return normalized


def validate_and_normalize_rows(rows):
    if not rows:
        return False, "Uploaded data is empty", [], []
    if len(rows) > 1000:
        return False, "Upload supports up to 1000 students", [], []

    cleaned_rows = [normalize_row_keys(row) for row in rows if row]
    if not cleaned_rows:
        return False, "Uploaded data is empty", [], []

    first_keys = list(cleaned_rows[0].keys())
    missing_required = REQUIRED_COLUMNS - set(first_keys)
    if missing_required:
        return False, "Required columns: StudentID, StudentName", [], []

    subject_columns = [key for key in first_keys if key not in REQUIRED_COLUMNS]
    if not subject_columns:
        return False, "At least one subject column is required", [], []

    normalized_rows = []
    expected_keys = set(first_keys)

    for row in cleaned_rows:
        if set(row.keys()) != expected_keys:
            return False, "All rows must have the same columns as the header", [], []

        student_id = str(row.get("StudentID", "")).strip()
        student_name = str(row.get("StudentName", "")).strip()
        if not student_id or not student_name:
            return False, "StudentID and StudentName cannot be empty", [], []

        subjects = {}
        for subject_name in subject_columns:
            raw_value = row.get(subject_name)
            try:
                numeric_value = round(float(str(raw_value).strip()), 2)
            except (TypeError, ValueError):
                return False, f"Column '{subject_name}' must contain numeric values only", [], []

            if numeric_value < 0 or numeric_value > 100:
                return False, f"Column '{subject_name}' must have values between 0 and 100", [], []

            subjects[subject_name] = numeric_value

        normalized_rows.append({"student_id": student_id, "name": student_name, "subjects": subjects})

    return True, "Data is valid", normalized_rows, subject_columns


def parse_rows_from_text(text_content):
    reader = csv.DictReader(io.StringIO(text_content.strip()))
    return [normalize_row_keys(row) for row in reader if row]


def parse_rows_from_uploaded_file(save_path, ext):
    if ext == ".json":
        with open(save_path, "r", encoding="utf-8") as file_obj:
            payload = json.load(file_obj)
        if isinstance(payload, dict) and isinstance(payload.get("students"), list):
            return payload["students"]
        if isinstance(payload, list):
            return payload
        raise ValueError("JSON must be a list of student objects or an object with a 'students' list")

    if ext == ".txt":
        with open(save_path, "r", encoding="utf-8") as file_obj:
            return parse_rows_from_text(file_obj.read())

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF parsing dependency missing. Install 'pypdf'.") from exc

        reader = PdfReader(save_path)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return parse_rows_from_text(text)

    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX parsing dependency missing. Install 'python-docx'.") from exc

        document = Document(save_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return parse_rows_from_text(text)

    if ext == ".doc":
        raise ValueError("Legacy .doc is not supported for parsing. Convert the file to .docx or .txt.")

    raise ValueError("Unsupported file type")


def compute_grade(average):
    for grade, threshold in GRADE_RULES:
        if average >= threshold:
            return grade
    return "F"


def is_pass(average):
    return average >= 50


def process_student_data(rows, class_id, class_name, uploaded_at):
    scored_rows = []
    for row in rows:
        subject_values = list(row["subjects"].values())
        total = round(sum(subject_values), 2)
        average = round(total / len(subject_values), 2)
        best_subject = max(row["subjects"], key=row["subjects"].get)
        grade = compute_grade(average)
        pass_status = "Pass" if is_pass(average) else "Fail"

        scored_rows.append(
            {
                "class_id": class_id,
                "class_name": class_name,
                "uploaded_at": uploaded_at,
                "student_id": row["student_id"],
                "name": row["name"],
                "subjects": row["subjects"],
                "total": total,
                "average": average,
                "best_subject": best_subject,
                "grade": grade,
                "pass_status": pass_status,
            }
        )

    ranked_rows = sorted(scored_rows, key=lambda item: (-item["total"], -item["average"], item["name"].lower()))
    for index, row in enumerate(ranked_rows, start=1):
        row["rank"] = index
    return ranked_rows


def get_latest_class_doc():
    return classes_collection.find_one({}, sort=[("uploaded_at", -1)])


def get_class_doc(class_id=None):
    if class_id:
        return classes_collection.find_one({"class_id": class_id})
    return get_latest_class_doc()


def get_students_for_class(class_id):
    return list(students_collection.find({"class_id": class_id}, {"_id": 0}))


def summarize_class(student_docs, class_doc):
    return {
        "class_id": class_doc["class_id"],
        "class_name": class_doc["class_name"],
        "uploaded_at": to_iso(class_doc["uploaded_at"]),
        "student_count": len(student_docs),
        "subject_count": len(class_doc.get("subject_columns", [])),
        "subjects": class_doc.get("subject_columns", []),
    }


def compute_dashboard_metrics(student_docs, class_doc):
    if not class_doc:
        return {
            "summary": {
                "class_id": None,
                "class_name": None,
                "uploaded_at": None,
                "total_students": 0,
                "class_average": 0,
                "subject_count": 0,
                "subjects": [],
                "topper": None,
            },
            "top_students": [],
            "top_rankings": [],
            "subject_averages": {},
            "pass_fail": {"pass_count": 0, "fail_count": 0},
            "grade_distribution": {"A": 0, "B": 0, "C": 0, "F": 0},
            "class_performance": {"labels": [], "totals": [], "averages": []},
            "insights": {
                "topper": None,
                "weakest_subject": None,
                "best_subject": None,
                "students_below_class_average": 0,
            },
        }

    subjects = class_doc.get("subject_columns", [])
    if not student_docs:
        return {
            "summary": {**summarize_class(student_docs, class_doc), "total_students": 0, "class_average": 0, "topper": None},
            "top_students": [],
            "top_rankings": [],
            "subject_averages": {subject: 0 for subject in subjects},
            "pass_fail": {"pass_count": 0, "fail_count": 0},
            "grade_distribution": {"A": 0, "B": 0, "C": 0, "F": 0},
            "class_performance": {"labels": [], "totals": [], "averages": []},
            "insights": {
                "topper": None,
                "weakest_subject": None,
                "best_subject": None,
                "students_below_class_average": 0,
            },
        }

    sorted_by_rank = sorted(student_docs, key=lambda item: item["rank"])
    class_average = round(sum(student["average"] for student in student_docs) / len(student_docs), 2)
    subject_averages = {
        subject: round(sum(student["subjects"].get(subject, 0) for student in student_docs) / len(student_docs), 2)
        for subject in subjects
    }

    grade_distribution = {"A": 0, "B": 0, "C": 0, "F": 0}
    pass_count = 0
    fail_count = 0
    for student in student_docs:
        grade_distribution[student["grade"]] = grade_distribution.get(student["grade"], 0) + 1
        if student["pass_status"] == "Pass":
            pass_count += 1
        else:
            fail_count += 1

    topper = sorted_by_rank[0]
    weakest_subject = min(subject_averages, key=subject_averages.get) if subject_averages else None
    best_subject = max(subject_averages, key=subject_averages.get) if subject_averages else None
    below_class_average = [student for student in student_docs if student["average"] < class_average]

    return {
        "summary": {
            **summarize_class(student_docs, class_doc),
            "total_students": len(student_docs),
            "class_average": class_average,
            "topper": {"name": topper["name"], "total": topper["total"], "average": topper["average"]},
        },
        "top_students": [
            {"name": student["name"], "total": student["total"], "rank": student["rank"], "average": student["average"]}
            for student in sorted_by_rank[:5]
        ],
        "top_rankings": [
            {"rank": student["rank"], "name": student["name"], "average": student["average"], "grade": student["grade"]}
            for student in sorted_by_rank[:5]
        ],
        "subject_averages": subject_averages,
        "pass_fail": {"pass_count": pass_count, "fail_count": fail_count},
        "grade_distribution": grade_distribution,
        "class_performance": {
            "labels": [student["name"] for student in sorted_by_rank[:20]],
            "totals": [student["total"] for student in sorted_by_rank[:20]],
            "averages": [student["average"] for student in sorted_by_rank[:20]],
        },
        "insights": {
            "topper": {"name": topper["name"], "total": topper["total"]},
            "weakest_subject": {"name": weakest_subject, "average": subject_averages.get(weakest_subject, 0)} if weakest_subject else None,
            "best_subject": {"name": best_subject, "average": subject_averages.get(best_subject, 0)} if best_subject else None,
            "students_below_class_average": len(below_class_average),
        },
    }


def get_upload_history(limit=8):
    history = []
    for class_doc in classes_collection.find({}, {"_id": 0}).sort("uploaded_at", -1).limit(limit):
        history.append(
            {
                "class_id": class_doc["class_id"],
                "class_name": class_doc["class_name"],
                "uploaded_at": to_iso(class_doc["uploaded_at"]),
                "student_count": class_doc.get("student_count", 0),
            }
        )
    return history


def list_classes_payload():
    classes = []
    for class_doc in classes_collection.find({}, {"_id": 0}).sort("uploaded_at", -1):
        classes.append(
            {
                "class_id": class_doc["class_id"],
                "class_name": class_doc["class_name"],
                "uploaded_at": to_iso(class_doc["uploaded_at"]),
                "student_count": class_doc.get("student_count", 0),
                "subject_count": len(class_doc.get("subject_columns", [])),
                "source_file": class_doc.get("source_file"),
            }
        )
    return classes


def build_student_response(student_doc, class_doc, student_docs, email=None):
    dashboard = compute_dashboard_metrics(student_docs, class_doc)
    class_average = dashboard["summary"]["class_average"]
    weak_subject = min(student_doc["subjects"], key=student_doc["subjects"].get) if student_doc["subjects"] else None
    performance_message = (
        f"You are above class average in {class_doc['class_name']}."
        if student_doc["average"] >= class_average
        else f"Needs improvement in {weak_subject}."
    )

    return {
        "student": {
            "student_id": student_doc["student_id"],
            "name": student_doc["name"],
            "email": email or "Not linked",
            "subjects": student_doc["subjects"],
            "total": student_doc["total"],
            "average": student_doc["average"],
            "rank": student_doc["rank"],
            "best_subject": student_doc["best_subject"],
            "weak_subject": weak_subject,
            "grade": student_doc["grade"],
            "pass_status": student_doc["pass_status"],
        },
        "class": {**summarize_class(student_docs, class_doc), "class_average": class_average, "subject_averages": dashboard["subject_averages"]},
        "comparison": {"student_average": student_doc["average"], "class_average": class_average},
        "message": performance_message,
    }


def resolve_student_by_name_or_email(query, class_id):
    class_doc = get_class_doc(class_id)
    if not class_doc:
        return None, None, None, None

    student_docs = get_students_for_class(class_doc["class_id"])
    if not student_docs:
        return None, class_doc, [], None

    query = str(query or "").strip()
    if not query:
        return None, class_doc, student_docs, None

    matched_user = users_collection.find_one(
        {
            "role": "student",
            "$or": [
                {"email": query.lower()},
                {"email": re.compile(re.escape(query), re.IGNORECASE)},
                {"name": re.compile(re.escape(query), re.IGNORECASE)},
            ],
        },
        {"_id": 0, "name": 1, "email": 1},
    )

    student_doc = None
    linked_email = None

    if matched_user:
        linked_email = matched_user.get("email")
        name_pattern = re.compile(f"^{re.escape(matched_user['name'])}$", re.IGNORECASE)
        student_doc = students_collection.find_one({"class_id": class_doc["class_id"], "name": name_pattern}, {"_id": 0})

    if not student_doc:
        exact_pattern = re.compile(f"^{re.escape(query)}$", re.IGNORECASE)
        student_doc = students_collection.find_one({"class_id": class_doc["class_id"], "name": exact_pattern}, {"_id": 0})

    if not student_doc:
        student_doc = students_collection.find_one(
            {"class_id": class_doc["class_id"], "name": re.compile(re.escape(query), re.IGNORECASE)},
            {"_id": 0},
        )

    if student_doc and not linked_email:
        linked_user = users_collection.find_one(
            {"role": "student", "name": re.compile(f"^{re.escape(student_doc['name'])}$", re.IGNORECASE)},
            {"_id": 0, "email": 1},
        )
        linked_email = linked_user.get("email") if linked_user else None

    return student_doc, class_doc, student_docs, linked_email


def get_student_classes(user_name):
    classes = []
    seen = set()
    class_docs = list(classes_collection.find({}, {"_id": 0}).sort("uploaded_at", -1))
    for class_doc in class_docs:
        student_doc = students_collection.find_one(
            {
                "class_id": class_doc["class_id"],
                "name": re.compile(f"^{re.escape(user_name.strip())}$", re.IGNORECASE),
            },
            {"_id": 0},
        )
        if student_doc and class_doc["class_id"] not in seen:
            seen.add(class_doc["class_id"])
            classes.append(
                {
                    "class_id": class_doc["class_id"],
                    "class_name": class_doc["class_name"],
                    "uploaded_at": to_iso(class_doc["uploaded_at"]),
                }
            )
    return classes


def get_filtered_students(class_id, search_text="", sort_key="rank", filter_key="all"):
    class_doc = get_class_doc(class_id)
    if not class_doc:
        return None, []

    students = get_students_for_class(class_doc["class_id"])
    search_text = str(search_text or "").strip().lower()

    if search_text:
        students = [
            student
            for student in students
            if search_text in student["name"].lower() or search_text in student["student_id"].lower()
        ]

    dashboard = compute_dashboard_metrics(get_students_for_class(class_doc["class_id"]), class_doc)
    class_average = dashboard["summary"]["class_average"]

    if filter_key == "top":
        students = [student for student in students if student["rank"] <= 5]
    elif filter_key == "average":
        students = [student for student in students if student["average"] >= class_average]
    elif filter_key == "weak":
        students = [student for student in students if student["average"] < class_average]
    elif filter_key == "fail":
        students = [student for student in students if student["pass_status"] == "Fail"]
    elif filter_key == "pass":
        students = [student for student in students if student["pass_status"] == "Pass"]

    sort_map = {
        "rank": lambda item: item["rank"],
        "average": lambda item: (-item["average"], item["rank"]),
        "total": lambda item: (-item["total"], item["rank"]),
        "name": lambda item: item["name"].lower(),
    }
    students = sorted(students, key=sort_map.get(sort_key, sort_map["rank"]))
    return class_doc, students


def export_excel_response(class_doc, student_docs):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ValueError("Excel export dependency missing. Install 'openpyxl'.") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    subjects = class_doc.get("subject_columns", [])
    headers = ["Student ID", "Student Name", *subjects, "Total", "Average", "Rank", "Grade", "Status"]
    sheet.append(headers)

    for student in sorted(student_docs, key=lambda item: item["rank"]):
        sheet.append(
            [
                student["student_id"],
                student["name"],
                *[student["subjects"].get(subject, 0) for subject in subjects],
                student["total"],
                student["average"],
                student["rank"],
                student["grade"],
                student["pass_status"],
            ]
        )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{class_doc['class_name'].replace(' ', '_').lower()}_report.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def export_pdf_response(class_doc, student_docs):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ValueError("PDF export dependency missing. Install 'reportlab'.") from exc

    dashboard = compute_dashboard_metrics(student_docs, class_doc)
    subjects = class_doc.get("subject_columns", [])
    rows = [["Name", *subjects, "Total", "Average", "Rank", "Grade", "Status"]]
    for student in sorted(student_docs, key=lambda item: item["rank"]):
        rows.append(
            [
                student["name"],
                *[student["subjects"].get(subject, 0) for subject in subjects],
                student["total"],
                student["average"],
                student["rank"],
                student["grade"],
                student["pass_status"],
            ]
        )

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"{class_doc['class_name']} Report", styles["Title"]),
        Spacer(1, 12),
        Paragraph(f"Uploaded: {to_iso(class_doc['uploaded_at'])}", styles["BodyText"]),
        Paragraph(f"Class Average: {dashboard['summary']['class_average']}", styles["BodyText"]),
        Paragraph(
            f"Topper: {dashboard['summary']['topper']['name']} ({dashboard['summary']['topper']['total']})"
            if dashboard["summary"]["topper"]
            else "Topper: N/A",
            styles["BodyText"],
        ),
        Spacer(1, 12),
    ]

    table = Table(rows, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16315f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cfd6e8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    output.seek(0)

    filename = f"{class_doc['class_name'].replace(' ', '_').lower()}_report.pdf"
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/pdf")


def export_student_pdf_response(class_doc, student_doc, class_average):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ValueError("PDF export dependency missing. Install 'reportlab'.") from exc

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()

    subject_rows = [["Subject", "Marks"]]
    for subject, marks in student_doc["subjects"].items():
        subject_rows.append([subject, marks])

    summary_rows = [
        ["Student Name", student_doc["name"]],
        ["Student ID", student_doc["student_id"]],
        ["Class", class_doc["class_name"]],
        ["Rank", student_doc["rank"]],
        ["Grade", student_doc["grade"]],
        ["Status", student_doc["pass_status"]],
        ["Total", student_doc["total"]],
        ["Average", student_doc["average"]],
        ["Class Average", class_average],
        ["Best Subject", student_doc["best_subject"]],
    ]

    summary_table = Table(summary_rows, colWidths=[140, 260])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eff5ff")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#17315d")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d4dcec")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ]
        )
    )

    subjects_table = Table(subject_rows, colWidths=[220, 120])
    subjects_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16315f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d4dcec")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fc")]),
            ]
        )
    )

    elements = [
        Paragraph("Student Marks Report", styles["Title"]),
        Spacer(1, 10),
        Paragraph(f"Generated for {student_doc['name']}", styles["Heading3"]),
        Paragraph(f"Class: {class_doc['class_name']}", styles["BodyText"]),
        Paragraph(f"Uploaded: {to_iso(class_doc['uploaded_at'])}", styles["BodyText"]),
        Spacer(1, 14),
        summary_table,
        Spacer(1, 16),
        Paragraph("Subject-wise Marks", styles["Heading3"]),
        Spacer(1, 8),
        subjects_table,
    ]

    doc.build(elements)
    output.seek(0)
    filename = f"{student_doc['name'].replace(' ', '_').lower()}_{class_doc['class_name'].replace(' ', '_').lower()}_marks.pdf"
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/pdf")


@app.route("/", methods=["GET"])
def health():
    return jsonify({"message": "Smart Student Data Visualization Portal backend is running"}), 200


@app.route("/register", methods=["POST"])
def register():
    data = request.get_json() or {}

    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", "")).strip()
    role = str(data.get("role", "")).strip().lower()

    if not all([name, email, password, role]):
        return jsonify({"error": "All fields are required: name, email, password, role"}), 400
    if role not in {"admin", "student"}:
        return jsonify({"error": "Role must be either 'admin' or 'student'"}), 400
    if not is_valid_email(email):
        return jsonify({"error": "Invalid email format"}), 400
    if len(password) < 6:
        return jsonify({"error": "Password must be at least 6 characters"}), 400
    if users_collection.find_one({"email": email}):
        return jsonify({"error": "Email already registered"}), 409

    hashed_password = bcrypt.generate_password_hash(password).decode("utf-8")
    users_collection.insert_one({"name": name, "email": email, "password": hashed_password, "role": role})
    return jsonify({"message": "Registration successful"}), 201


@app.route("/login", methods=["POST"])
def login():
    data = request.get_json() or {}

    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", "")).strip()
    if not email or not password:
        return jsonify({"error": "Email and password are required"}), 400

    user = users_collection.find_one({"email": email})
    if not user or not bcrypt.check_password_hash(user["password"], password):
        return jsonify({"error": "Invalid email or password"}), 401

    token = create_access_token(identity=user["email"], additional_claims={"role": user["role"], "name": user["name"]})
    return jsonify(
        {
            "message": "Login successful",
            "token": token,
            "user": {"name": user["name"], "email": user["email"], "role": user["role"]},
        }
    ), 200


@app.route("/upload_data", methods=["POST"])
@role_required("admin")
def upload_data():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded. Use form-data key 'file'"}), 400

    uploaded_file = request.files["file"]
    if uploaded_file.filename == "":
        return jsonify({"error": "Empty filename"}), 400

    _, ext = os.path.splitext(uploaded_file.filename.lower())
    if ext not in ALLOWED_UPLOAD_EXTENSIONS:
        return jsonify({"error": "Only .json, .txt, .pdf, .docx, .doc files are allowed"}), 400

    class_name = normalize_class_name(request.form.get("class_name"))
    class_id = build_class_id()
    uploaded_at = utc_now()

    save_path = os.path.join(UPLOAD_DIR, f"{class_id}_{uploaded_file.filename}")
    uploaded_file.save(save_path)

    try:
        parsed_rows = parse_rows_from_uploaded_file(save_path, ext)
        is_valid, message, normalized_rows, subject_columns = validate_and_normalize_rows(parsed_rows)
        if not is_valid:
            return jsonify({"error": message}), 400
    except Exception as exc:
        return jsonify({"error": f"Invalid uploaded file: {str(exc)}"}), 400

    student_docs = process_student_data(normalized_rows, class_id, class_name, uploaded_at)
    students_collection.insert_many(student_docs)
    classes_collection.insert_one(
        {
            "class_id": class_id,
            "class_name": class_name,
            "uploaded_at": uploaded_at,
            "subject_columns": subject_columns,
            "student_count": len(student_docs),
            "source_file": uploaded_file.filename,
        }
    )

    dashboard_data = compute_dashboard_metrics(student_docs, get_class_doc(class_id))
    return jsonify(
        {
            "message": "File uploaded and processed successfully",
            "class_id": class_id,
            "class_name": class_name,
            "students_processed": len(student_docs),
            "dashboard_preview": dashboard_data,
        }
    ), 200


@app.route("/admin/classes", methods=["GET"])
@role_required("admin")
def admin_classes():
    return jsonify({"classes": list_classes_payload()}), 200


@app.route("/admin/dashboard", methods=["GET"])
@role_required("admin")
def admin_dashboard():
    class_id = request.args.get("class_id")
    class_doc = get_class_doc(class_id)
    if not class_doc:
        return jsonify({"error": "No uploaded class data found"}), 404

    student_docs = get_students_for_class(class_doc["class_id"])
    payload = compute_dashboard_metrics(student_docs, class_doc)
    payload["available_classes"] = list_classes_payload()
    payload["upload_history"] = get_upload_history()
    return jsonify(payload), 200


@app.route("/admin/student-search", methods=["GET"])
@role_required("admin")
def admin_student_search():
    query = str(request.args.get("q", "")).strip()
    class_id = request.args.get("class_id")
    if not query:
        return jsonify({"error": "Search query is required"}), 400

    student_doc, class_doc, student_docs, linked_email = resolve_student_by_name_or_email(query, class_id)
    if not class_doc:
        return jsonify({"error": "No uploaded class data found"}), 404
    if not student_doc:
        return jsonify({"error": "No student found for the provided name or email"}), 404

    response = build_student_response(student_doc, class_doc, student_docs, linked_email)
    response["available_classes"] = list_classes_payload()
    return jsonify(response), 200


@app.route("/admin/students", methods=["GET"])
@role_required("admin")
def admin_students():
    class_id = request.args.get("class_id")
    search_text = request.args.get("search", "")
    sort_key = request.args.get("sort", "rank")
    filter_key = request.args.get("filter", "all")

    class_doc, students = get_filtered_students(class_id, search_text, sort_key, filter_key)
    if not class_doc:
        return jsonify({"error": "No uploaded class data found"}), 404

    return jsonify(
        {
            "class": {
                "class_id": class_doc["class_id"],
                "class_name": class_doc["class_name"],
                "uploaded_at": to_iso(class_doc["uploaded_at"]),
                "subjects": class_doc.get("subject_columns", []),
            },
            "students": [
                {
                    "student_id": student["student_id"],
                    "name": student["name"],
                    "total": student["total"],
                    "average": student["average"],
                    "rank": student["rank"],
                    "grade": student["grade"],
                    "pass_status": student["pass_status"],
                    "best_subject": student["best_subject"],
                }
                for student in students
            ],
        }
    ), 200


@app.route("/admin/export/excel", methods=["GET"])
@role_required("admin")
def admin_export_excel():
    class_id = request.args.get("class_id")
    class_doc = get_class_doc(class_id)
    if not class_doc:
        return jsonify({"error": "No uploaded class data found"}), 404

    student_docs = get_students_for_class(class_doc["class_id"])
    try:
        return export_excel_response(class_doc, student_docs)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/admin/export/pdf", methods=["GET"])
@role_required("admin")
def admin_export_pdf():
    class_id = request.args.get("class_id")
    class_doc = get_class_doc(class_id)
    if not class_doc:
        return jsonify({"error": "No uploaded class data found"}), 404

    student_docs = get_students_for_class(class_doc["class_id"])
    try:
        return export_pdf_response(class_doc, student_docs)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/student/dashboard", methods=["GET"])
@role_required("student")
def student_dashboard():
    identity = get_jwt_identity()
    email = str(identity or "").strip().lower()
    class_id = request.args.get("class_id")

    user = users_collection.find_one({"email": email}, {"_id": 0, "name": 1, "email": 1})
    if not user:
        return jsonify({"error": "User not found"}), 404

    available_classes = get_student_classes(user["name"])
    if not available_classes:
        return jsonify({"error": "No uploaded student record found for this user"}), 404

    if not class_id:
        class_id = available_classes[0]["class_id"]

    class_doc = get_class_doc(class_id)
    if not class_doc:
        return jsonify({"error": "Selected class not found"}), 404

    student_doc = students_collection.find_one(
        {"class_id": class_doc["class_id"], "name": re.compile(f'^{re.escape(user["name"])}$', re.IGNORECASE)},
        {"_id": 0},
    )
    if not student_doc:
        return jsonify({"error": "No uploaded student record found for this class"}), 404

    student_docs = get_students_for_class(class_doc["class_id"])
    payload = build_student_response(student_doc, class_doc, student_docs, user.get("email"))
    payload["available_classes"] = available_classes
    return jsonify(payload), 200


@app.route("/student/export/pdf", methods=["GET"])
@role_required("student")
def student_export_pdf():
    identity = get_jwt_identity()
    email = str(identity or "").strip().lower()
    class_id = request.args.get("class_id")

    user = users_collection.find_one({"email": email}, {"_id": 0, "name": 1, "email": 1})
    if not user:
        return jsonify({"error": "User not found"}), 404

    available_classes = get_student_classes(user["name"])
    if not available_classes:
        return jsonify({"error": "No uploaded student record found for this user"}), 404

    if not class_id:
        class_id = available_classes[0]["class_id"]

    class_doc = get_class_doc(class_id)
    if not class_doc:
        return jsonify({"error": "Selected class not found"}), 404

    student_doc = students_collection.find_one(
        {"class_id": class_doc["class_id"], "name": re.compile(f'^{re.escape(user["name"])}$', re.IGNORECASE)},
        {"_id": 0},
    )
    if not student_doc:
        return jsonify({"error": "No uploaded student record found for this class"}), 404

    student_docs = get_students_for_class(class_doc["class_id"])
    dashboard = compute_dashboard_metrics(student_docs, class_doc)

    try:
        return export_student_pdf_response(class_doc, student_doc, dashboard["summary"]["class_average"])
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 500


@jwt.unauthorized_loader
def unauthorized_callback(error):
    return jsonify({"error": "Missing or invalid Authorization header", "details": error}), 401


@jwt.invalid_token_loader
def invalid_token_callback(error):
    return jsonify({"error": "Invalid token", "details": error}), 422


@jwt.expired_token_loader
def expired_token_callback(jwt_header, jwt_payload):
    return jsonify({"error": "Token has expired"}), 401


if __name__ == "__main__":
    app.run(debug=True)
